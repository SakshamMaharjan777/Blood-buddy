# -*- coding: utf-8 -*-
"""
Builds the BloodBuddy Final Project Report (DOCX) using the reference report's
formatting: Arial family, A4, same front-matter order, chapter/section hierarchy,
captioned tables, bracketed figure placeholders, centered page numbers.

Run:  python tools/report/build_report.py
Then (page numbers + TOC + PDF):  python tools/report/finalize_word.py
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import bb_content as C

BODY = "Arial"          # reference uses ArialMT
CODE_FONT = "Courier New"   # reference uses CourierNewPSMT
CODE_BG = "F2F2F2"          # light grey behind code blocks
ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")

# Figures for which we have real image assets: caption text -> (image file, width),
# or a LIST of (file, width) pairs when one figure shows several screens.
# Diagrams were extracted from the proposal DOCX; the shots/ tree holds screenshots
# captured from the running application (see PROGRESS.md session #26).
FIG_IMAGES = {
    "Figure 3-1: N-Tier Architecture of BloodBuddy": ("fig3_1_architecture.png", Cm(14)),
    "Figure 3-3: Entity Relationship Diagram": ("fig3_3_erd.png", Cm(14)),
    "Figure 3-4: Class Diagram of Core Entities": ("fig3_2_class_diagram.png", Cm(14)),
    "Figure 6-1: Project Gantt Chart": ("fig6_1_gantt.png", Cm(16)),
    # Drawn by build_diagrams.py (PIL-only, reproducible: they carry no placeholder)
    "Figure 1-1: BloodBuddy Request Lifecycle": ("fig1_1_lifecycle.png", Cm(16)),
    "Figure 3-2: Use Case Diagram": ("fig3_2_use_case.png", Cm(16)),
    "Figure 3-5: Level-1 Data Flow Diagram": ("fig3_5_dfd.png", Cm(16)),
    "Figure 4-1: Project File Structure": ("fig4_1_structure.png", Cm(16)),
    # Screenshots (individual captures, or the grouped montages composed by
    # build_montages.py where one figure number asks for more than one screen)
    "Figure 4-2: Donor Search with Filters and Pagination": ("shots/fig4_2_donor_search.png", Cm(15)),
    "Figure 4-3: Blood Request Form with Inline Validation": ("shots/fig4_3_request_forms.png", Cm(13.5)),
    "Figure 4-4: Profile Photo Upload (Base64 Preview Round-Trip)": ("shots/fig4_4_donor_profile.png", Cm(14)),
    "Figure 4-5: Administrator Dashboard with Aggregate Analytics": ("shots/fig4_5_admin_panels.png", Cm(16)),
    "Figure 5-1: Automated Verification Reports (268/268 · 92/92 · 42/42)": ("shots/fig5_1_harness.png", Cm(13.5)),
}

# Appendix figures, keyed by appendix title: the appendix renders its prose and any
# table first, then these images, each with its own caption.
APPENDIX_IMAGES = {
    "Appendix E: QA Evidence and Reproduction Steps": [
        ("shots/fig5_1_harness.png", Cm(13.5),
         "Figure E-1: Automated Verification Reports (reproduced from Figure 5-1)"),
    ],
    "Appendix F: Additional Screenshots": [
        ("shots/appF_landing.png", Cm(13), "Figure F-1: Landing Page"),
        ("shots/appF_auth_pair.png", Cm(15),
         "Figure F-2: Authentication Screens (Login and Registration)"),
        ("shots/appF_donor_dashboard.png", Cm(13), "Figure F-3: Donor Dashboard"),
        ("shots/appF_hospital_requests.png", Cm(13), "Figure F-4: Hospital Request Queue"),
        ("shots/appF_legal_row.png", Cm(15),
         "Figure F-5: Terms, Privacy Policy and Branded 404"),
    ],
}

document = Document()

# ------------------------------------------------------------------ page setup
sec = document.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)   # A4 (reference MediaBox 595x842pt)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
sec.left_margin, sec.right_margin = Cm(2.4), Cm(2.4)
sec.different_first_page_header_footer = True   # cover shows no page number (as in reference)

styles = document.styles

def _force_font(style, name):
    """Set the font AND clear Word theme-font overrides, which otherwise win."""
    style.font.name = name
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rFonts.attrib.pop(qn("w:" + attr), None)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)

st = styles["Normal"]
_force_font(st, BODY)
st.font.size = Pt(11)
st._element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15

h1 = styles["Heading 1"]
_force_font(h1, BODY)
h1.font.size = Pt(15)
h1.font.bold = True
h1.font.color.rgb = RGBColor(0, 0, 0)
h1.paragraph_format.space_before = Pt(6)
h1.paragraph_format.space_after = Pt(12)
h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

h2 = styles["Heading 2"]
_force_font(h2, BODY)
h2.font.size = Pt(13)
h2.font.bold = True
h2.font.color.rgb = RGBColor(0, 0, 0)
h2.paragraph_format.space_before = Pt(12)
h2.paragraph_format.space_after = Pt(6)
h2.paragraph_format.keep_with_next = True

h3 = styles["Heading 3"]
_force_font(h3, BODY)
h3.font.size = Pt(11.5)
h3.font.bold = True
h3.font.color.rgb = RGBColor(0, 0, 0)
h3.paragraph_format.space_before = Pt(10)
h3.paragraph_format.space_after = Pt(4)
h3.paragraph_format.keep_with_next = True

cap = styles.add_style("BB Caption", 1)
_force_font(cap, BODY)
cap.font.size = Pt(10.5)
cap.font.bold = True
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_before = Pt(8)
cap.paragraph_format.space_after = Pt(4)
cap.paragraph_format.keep_with_next = True

# --------------------------------------------------------------- helpers

def _shade(cell, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _add_field(par, instr):
    r1, r2, r3 = par.add_run(), par.add_run(), par.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = " "
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r1._r.append(b); r2._r.append(i); r3._r.append(s); r3._r.append(t); r3._r.append(e)


def para(text, align=None, space_after=6, bold=False, italic=False,
         size=None, font=None, line=None):
    p = document.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if line:
        p.paragraph_format.line_spacing = line
    r = p.add_run(text)
    r.font.bold = bold
    r.font.italic = italic
    if size:
        r.font.size = Pt(size)
    if font:
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return p


def bullet(text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def heading(text, level=1):
    h = document.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = BODY
        r.font.color.rgb = RGBColor(0, 0, 0)
    return h


def page_break():
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def caption(text):
    p = document.add_paragraph(style="BB Caption")
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.name = BODY
    return p


def add_table(spec):
    caption(spec["caption"])
    t = document.add_table(rows=1, cols=len(spec["headers"]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(spec["headers"]):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(htxt)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.name = BODY
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        _shade(c, "D9D9D9")
    for row in spec["rows"]:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(val)
            r.font.size = Pt(9.5)
            r.font.name = BODY
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
    para("", space_after=4)
    return t


def add_image(rel_path, width):
    """Insert one centered image. Returns False (loudly) if the asset is missing."""
    path = os.path.join(ASSETS, rel_path)
    if not os.path.exists(path):
        print("WARNING: missing image asset:", path)
        return False
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=width)
    return True


def _fig_specs(entry):
    """A FIG_IMAGES value is one (file, width) pair, or a list of them."""
    if isinstance(entry[0], str):
        return [entry]
    return list(entry)


def add_figure(caption_text, placeholder_lines):
    img = FIG_IMAGES.get(caption_text)
    if img:
        for rel_path, width in _fig_specs(img):
            add_image(rel_path, width)
    for ln in placeholder_lines:
        p = para(ln, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=9.5,
                 space_after=0)
        p.paragraph_format.keep_with_next = True
    caption(caption_text)
    para("", space_after=6)


def add_code(code_text):
    for ln in code_text.split("\n"):
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.6)
        r = p.add_run(ln if ln else " ")
        r.font.name = CODE_FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
        r.font.size = Pt(8.5)
        _shade_p(p, CODE_BG)
    para("", space_after=4)


def _shade_p(p, hex_fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def render_blocks(blocks):
    for blk in blocks:
        kind = blk[0]
        if kind == "h1":
            heading(blk[1], 1)
        elif kind == "h2":
            heading(blk[1], 2)
        elif kind == "h3":
            heading(blk[1], 3)
        elif kind == "p":
            para(blk[1], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif kind == "b":
            bullet(blk[1])
        elif kind == "kv":
            for k, v in blk[1]:
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Cm(0.6)
                r1 = p.add_run(k + "   ")
                r1.font.bold = True
                r1.font.name = CODE_FONT
                r1.font.size = Pt(9)
                r2 = p.add_run(v)
                r2.font.name = CODE_FONT
                r2.font.size = Pt(9)
        elif kind == "code":
            add_code(blk[1])
        elif kind == "table":
            spec = blk[1]
            if isinstance(spec, str):
                spec = getattr(C, spec)
            add_table(spec)
        elif kind == "fig":
            add_figure(blk[1], blk[2])


# ------------------------------------------------------------------ footers
def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_field(p, "PAGE")


add_page_number_footer(sec)

# ================================================================== FRONT MATTER
def front_page():
    # University logo at the top of the cover (same position as the reference report)
    logo = os.path.join(ASSETS, "cover_logo.png")
    if os.path.exists(logo):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(16)
        p.add_run().add_picture(logo, width=Cm(4))
    para("", space_after=0)
    para(C.TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16,
         space_after=20, line=1.25)
    para("FINAL PROJECT REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True,
         size=14, space_after=4)
    para(C.MODULE_LINE, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13,
         space_after=30)
    para("By:", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=4)
    para(C.STUDENT, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13,
         space_after=26)
    para(C.PROGRAMME, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12,
         space_after=8, line=1.25)
    para(C.DEPT, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, space_after=4)
    para(C.UNIVERSITY, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12,
         space_after=30)
    para(C.SUPERVISOR, align=WD_ALIGN_PARAGRAPH.CENTER, size=11.5, space_after=8)
    para(C.SUBMISSION_DATE, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11.5)
    page_break()


def centre_title(text):
    # Heading 1 so front-matter entries appear in the TOC (as in the reference report)
    h = heading(text, 1)
    for r in h.runs:
        r.font.size = Pt(14)
    return h


def disclaimer():
    centre_title("DISCLAIMER")
    for para_text in C.DISCLAIMER:
        para(para_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    page_break()


def acknowledgements():
    centre_title("ACKNOWLEDGEMENTS")
    for para_text in C.ACKNOWLEDGEMENTS:
        para(para_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    page_break()


def abstract():
    centre_title("ABSTRACT")
    for para_text in C.ABSTRACT:
        para(para_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(C.KEYWORDS, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=True, space_after=0)
    page_break()


def toc_placeholder():
    centre_title("TABLE OF CONTENTS")
    p = document.add_paragraph()
    _add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    page_break()


def simple_list_page(title, entries):
    centre_title(title)
    for e in entries:
        para(e, space_after=4)
    page_break()


def abbreviations():
    centre_title("LIST OF ABBREVIATIONS")
    t = document.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ab, meaning in C.ABBREVIATIONS:
        cells = t.add_row().cells
        r = cells[0].paragraphs[0].add_run(ab); r.font.bold = True
        r.font.size = Pt(10); r.font.name = BODY
        r2 = cells[1].paragraphs[0].add_run(meaning)
        r2.font.size = Pt(10); r2.font.name = BODY
    page_break()


front_page()
disclaimer()
acknowledgements()
abstract()
toc_placeholder()
simple_list_page("LIST OF TABLES", C.LIST_OF_TABLES)
simple_list_page("LIST OF FIGURES", C.LIST_OF_FIGURES)
abbreviations()

# ================================================================== CHAPTERS
# Chapter 2: splice the two comparison tables into section 2.3 (before 2.4)
CH2 = []
for blk in C.CH2:
    if blk[0] == "h2" and blk[1].startswith("2.4"):
        CH2.append(("table", C.TABLE_2_1))
        CH2.append(("table", C.TABLE_2_2))
    CH2.append(blk)

for blocks in (C.CH1, CH2, C.CH3_1, C.CH4, C.CH5, C.CH6, C.CH7):
    render_blocks(blocks)
    page_break()

# ================================================================== REFERENCES
heading("REFERENCES", 1)
for ref in C.REFERENCES:
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.space_after = Pt(8)
    p.add_run(ref)
page_break()

# ================================================================== APPENDICES
heading("APPENDICES", 1)
for title, paras_, table in C.APPENDICES:
    heading(title, 2)
    for t in paras_:
        para(t, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if table:
        add_table(table)
    for rel_path, width, cap_text in APPENDIX_IMAGES.get(title, []):
        if add_image(rel_path, width):
            caption(cap_text)
            para("", space_after=6)
    para("", space_after=6)

# ------------------------------------------------------------------ save
out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "..", "..", "docs", "BloodBuddy_Final_Report.docx")
if len(sys.argv) > 1:          # optional override: python build_report.py <output.docx>
    out = sys.argv[1]
document.save(os.path.normpath(out))
print("Saved:", os.path.normpath(out))
